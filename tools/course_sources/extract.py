"""Bounded private extraction/OCR pilot. No network or source modifications."""
import argparse, hashlib, json, re, time, zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET
from inventory import COURSE

def digest(path):
    h=hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda:stream.read(1024*1024),b''):h.update(block)
    return h.hexdigest()

def chunk_text(text,limit=1400,overlap=180):
    chunks=[];start=0
    while start<len(text):
        end=min(start+limit,len(text))
        chunks.append({'start':start,'end':end,'text':text[start:end]})
        if end==len(text):break
        start=end-overlap
    return chunks

def ocr_image(engine,path):
    rows,elapsed=engine(str(path))
    rows=rows or []
    return {'method':'rapidocr-onnxruntime-1.4.4','text':'\n'.join(r[1] for r in rows),
      'lines':[{'text':r[1],'confidence':float(r[2]),'polygon':[[float(v) for v in point] for point in r[0]]} for r in rows],
      'needsReview':True,'qualityNote':'OCR scores are model scores, not a guarantee of exact words, formulae or MCQ option order.'}

def xml_text(path):
    with zipfile.ZipFile(path) as z:
        names=[n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')]
        return [{'page':i+1,'method':'pptx-slide-text','text':'\n'.join(e.text for e in ET.fromstring(z.read(name)).iter() if e.tag.endswith('}t') and e.text),'needsReview':True} for i,name in enumerate(sorted(names,key=lambda n:int(re.search(r'slide(\d+)',n).group(1))))]

def extract_source(record,out,max_pages,engine=None):
    path=Path(record['path']);content_hash=digest(path);source_id=content_hash[:24]
    result={'sourceId':source_id,'contentSha256':content_hash,'inventoryId':record['id'],'originalPath':str(path),
      'kind':record['kind'],'courseCodes':record['courseCodes'],'privacy':'private-local','rights':record['rights'],
      'sourceInstructionsTrusted':False,'processedAtUnix':time.time(),'pages':[]}
    ext=path.suffix.lower()
    with path.open('rb') as source:signature=source.read(8)
    if ext=='.pdf' and signature.startswith(b'PK'):
        with zipfile.ZipFile(path) as archive:
            if 'word/document.xml' in archive.namelist():
                ext='.docx';result['formatWarning']='Filename says PDF; package is DOCX. Parsed content type without changing original.'
    if ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        if reader.is_encrypted:
            if not reader.decrypt(''):raise ValueError('Encrypted PDF requires user-supplied access; no bypass attempted')
        result['pageCount']=len(reader.pages)
        for i,page in enumerate(reader.pages[:max_pages]):
            text=page.extract_text() or '';entry={'page':i+1,'method':'pypdf-text','text':text,'needsReview':True,'ocrSuggested':len(text.strip())<100}
            if entry['ocrSuggested'] and engine:
                import pypdfium2 as pdfium
                with pdfium.PdfDocument(path) as rendered:
                    p=rendered[i];bitmap=p.render(scale=1.5);image=bitmap.to_pil();image_path=out/'renders'/f'{source_id}-p{i+1}.png'
                    image_path.parent.mkdir(exist_ok=True);image.save(image_path);bitmap.close();p.close()
                    entry.update(ocr_image(engine,image_path));entry['renderedPage']=str(image_path)
            result['pages'].append(entry)
    elif ext=='.docx':
        from docx import Document
        doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)
        tables=[[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]
        text+='\n'+'\n'.join(' | '.join(row) for table in tables for row in table)
        result['pages']=[{'page':None,'locator':'document paragraphs and tables; no reliable page numbers','method':'python-docx','text':text,'tables':tables,'needsReview':True}]
    elif ext=='.pptx':result['pages']=xml_text(path)[:max_pages]
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False)
        for sheet in book.worksheets[:max_pages]:
            rows=[[str(v) if v is not None else '' for v in row] for row in sheet.iter_rows(max_row=100, max_col=20,values_only=True)]
            result['pages'].append({'page':None,'sheet':sheet.title,'method':'openpyxl-read-only','text':'\n'.join(' | '.join(r) for r in rows),'needsReview':True,'formulasExecuted':False})
        book.close()
    elif ext in {'.txt','.md'}:result['pages']=[{'page':None,'method':'utf8-text','text':path.read_text(encoding='utf-8',errors='replace')[:60000],'needsReview':True}]
    elif engine:result['pages']=[{'page':1,**ocr_image(engine,path)}]
    else:raise ValueError('Image OCR engine unavailable')
    result['chunks']=[]
    for page in result['pages']:
        for n,chunk in enumerate(chunk_text(page['text'])):
            chunk.update({'chunkId':f"{source_id}:{page.get('page') or page.get('sheet') or 'document'}:{n}",
                'sourceId':source_id,'page':page.get('page'),'method':page['method'],'needsReview':True})
            result['chunks'].append(chunk)
    (out/'extracted').mkdir(exist_ok=True)
    (out/'extracted'/f'{source_id}.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result

def pick_sources(records,limit):
    # Diversity over codes and formats; priority is not a correctness judgement.
    candidates=[r for r in records if r['extension'] in {'.pdf','.docx','.pptx','.xlsx','.txt','.md'} and 500<r['bytes']<64*1024*1024]
    candidates.sort(key=lambda r:('student-work' in r['kind'],r['extension']!='.pdf',r['bytes']))
    selected=[];seen_codes=set();seen_format=set()
    for r in candidates:
        if len(selected)>=limit:break
        new_codes=set(r['courseCodes'])-seen_codes
        if new_codes or r['extension'] not in seen_format:
            selected.append(r);seen_codes.update(r['courseCodes']);seen_format.add(r['extension'])
    return selected

def run(args):
    out=args.private;inventory=json.loads((out/'inventory.json').read_text(encoding='utf-8'))
    engine=None
    if args.ocr:
        from rapidocr_onnxruntime import RapidOCR
        engine=RapidOCR(intra_op_num_threads=2,inter_op_num_threads=1)
    selected=pick_sources(inventory['sources'],args.max_files);summaries=[];errors=[]
    for record in selected:
        try:
            r=extract_source(record,out,args.pages,engine)
            summaries.append({'id':r['sourceId'],'courseCodes':r['courseCodes'],'extension':Path(record['path']).suffix,'pagesProcessed':len(r['pages']),'methods':list(set(p['method'] for p in r['pages'])),'characters':sum(len(p['text']) for p in r['pages']),'chunks':len(r['chunks']),'reviewStatus':'not-yet-validated'})
        except Exception as e:errors.append({'inventoryId':record['id'],'error':type(e).__name__+': '+str(e)[:240]})
    screenshots=[]
    if engine and args.screenshots:
        images=inventory['screenshots'];count=min(args.screenshots,len(images))
        positions=sorted(set(round(i*(len(images)-1)/max(1,count-1)) for i in range(count)))
        for pos in positions:
            r=images[pos]
            try:
                result=ocr_image(engine,Path(r['path']));text=result['text'];codes=sorted(set((a+b).upper() for a,b in COURSE.findall(text)))
                candidate=bool(codes or re.search(r'my\s*unisa|myModules',text,re.I) or (re.search(r'question\s*\d|select.*answer|multiple.choice',text,re.I) and re.search(r'SQL|database|econom|account|audit|management|probability|statistics',text,re.I)))
                item={'id':r['id'],'courseCodes':codes,'courseworkCandidate':candidate,'characters':len(text),'reviewStatus':'OCR candidate only; human review needed'}
                if candidate:
                    (out/'screenshot-ocr').mkdir(exist_ok=True)
                    (out/'screenshot-ocr'/f"{r['id']}.json").write_text(json.dumps({'path':r['path'],**item,**result},ensure_ascii=False,indent=2),encoding='utf-8')
                screenshots.append(item)
            except Exception as e:errors.append({'imageId':r['id'],'error':type(e).__name__+': '+str(e)[:240]})
    summary={'schemaVersion':1,'documentsAttempted':len(selected),'documentsExtracted':len(summaries),'documentPagesOrSections':sum(r['pagesProcessed'] for r in summaries),
      'extractedCharacters':sum(r['characters'] for r in summaries),'chunks':sum(r['chunks'] for r in summaries),'methods':dict(Counter(m for r in summaries for m in r['methods'])),
      'screenshotsSampled':len(screenshots),'screenshotCourseworkCandidates':sum(r['courseworkCandidate'] for r in screenshots),'errors':errors,
      'sources':summaries,'screenshotSample':screenshots,'cloudUploads':0,'originalsModified':False,'quality':'Extraction/OCR completed does not imply correctness or an answer key.'}
    (out/'extraction-summary.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in summary.items() if k not in ['sources','screenshotSample','errors']}|{'errorCount':len(errors)}))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--private',type=Path,default=Path('data/private/course-sources'))
    parser.add_argument('--max-files',type=int,default=24);parser.add_argument('--pages',type=int,default=2);parser.add_argument('--screenshots',type=int,default=24);parser.add_argument('--ocr',action='store_true')
    args=parser.parse_args()
    if not 1<=args.max_files<=50 or not 1<=args.pages<=10 or not 0<=args.screenshots<=100:parser.error('Bounded pilot limits: 1–50 files, 1–10 pages, 0–100 screenshots.')
    run(args)
